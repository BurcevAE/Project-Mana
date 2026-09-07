"""
mana.apps.documents — .docx and .xlsx without Word or Excel installed.

Office is not on this machine, and that turned out to be the better
constraint rather than a limitation to work around. python-docx and
openpyxl read and write the files themselves, so the same code works on a
machine with Office, on one without, and on a server -- and it is
testable here, which COM automation of an absent application is not.

What this cannot do, stated because the gap is not obvious
----------------------------------------------------------
These libraries manipulate the file, not the application. There is no
formula recalculation (openpyxl reads either the stored formula or the
value cached by whatever last opened the file, never a freshly computed
one), no printing, no PDF export, no interaction with a document a person
currently has open. Those need the application, and the application is
where the COM path would go.

`read_xlsx` therefore reports which of the two it gave you. A number that
was cached by Excel in 2019 and a number computed now are not the same
fact, and a reader who cannot tell them apart will eventually treat one
as the other.
"""
from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence

from . import require

#: Component version -- see mana/version.py for the bump conventions.
__version__ = "1.0"

#: Reading a whole workbook into a prompt is how a context window gets
#: spent on nothing. Callers that genuinely want everything pass their own
#: limit and say so.
DEFAULT_MAX_ROWS = 500


def read_docx(path: str) -> Dict[str, Any]:
    """Text, headings and tables of a Word document.

    Tables are returned separately from the paragraph flow rather than
    flattened into it: a table read as running text loses which cell
    belonged to which column, and that is usually the only information the
    table carried.
    """
    require("docx")
    import docx

    document = docx.Document(str(Path(path)))
    paragraphs, headings = [], []
    for paragraph in document.paragraphs:
        text = (paragraph.text or "").strip()
        if not text:
            continue
        style = (paragraph.style.name if paragraph.style else "") or ""
        paragraphs.append({"text": text, "style": style})
        if style.lower().startswith(("heading", "заголовок", "title")):
            headings.append(text)

    tables = []
    for table in document.tables:
        tables.append([[(cell.text or "").strip() for cell in row.cells]
                       for row in table.rows])

    return {
        "path": str(Path(path).resolve()),
        "paragraphs": paragraphs,
        "headings": headings,
        "tables": tables,
        "text": "\n".join(p["text"] for p in paragraphs),
    }


def write_docx(path: str, blocks: Sequence[Dict[str, Any]],
               overwrite: bool = False) -> Dict[str, Any]:
    """Create a Word document from a list of blocks.

    Each block is {"kind": "heading"|"text"|"table", ...}. Refuses to
    overwrite an existing file unless told to: this runs on behalf of an
    agent, and an agent that silently replaces a document the user spent
    an afternoon on has done something worse than failing.
    """
    require("docx")
    import docx

    target = Path(path)
    if target.exists() and not overwrite:
        raise FileExistsError(
            f"{target} уже существует; передайте overwrite=True, если заменить "
            f"его — это и есть намерение")

    document = docx.Document()
    for block in blocks:
        kind = str(block.get("kind", "text"))
        if kind == "heading":
            document.add_heading(str(block.get("text", "")),
                                 level=int(block.get("level", 1)))
        elif kind == "table":
            rows = block.get("rows") or []
            if not rows:
                continue
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = str(block.get("style", "Table Grid"))
            for r, row in enumerate(rows):
                for c, value in enumerate(row):
                    table.cell(r, c).text = "" if value is None else str(value)
        else:
            document.add_paragraph(str(block.get("text", "")))

    target.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(target))
    return {"path": str(target.resolve()), "blocks": len(blocks),
            "bytes": target.stat().st_size}


def read_xlsx(path: str, sheet: str = "", max_rows: int = DEFAULT_MAX_ROWS,
              formulas: bool = False) -> Dict[str, Any]:
    """Rows of a spreadsheet, and an honest note about what the values are.

    `formulas=False` returns the values cached in the file by whichever
    program last saved it; `formulas=True` returns the formula text. There
    is no third option that computes them, because computing them needs
    Excel. `values_are` in the result says which of the two you got, and
    `truncated` says whether the sheet was longer than `max_rows` -- a
    silently shortened table is a wrong table.
    """
    require("xlsx")
    import openpyxl

    book = openpyxl.load_workbook(str(Path(path)), data_only=not formulas,
                                  read_only=True)
    try:
        names = list(book.sheetnames)
        if sheet and sheet not in names:
            raise KeyError(f"листа {sheet!r} нет; есть: {', '.join(names)}")
        worksheet = book[sheet] if sheet else book[names[0]]

        rows: List[List[Any]] = []
        truncated = False
        for index, row in enumerate(worksheet.iter_rows(values_only=True)):
            if index >= max_rows:
                truncated = True
                break
            rows.append(list(row))

        return {
            "path": str(Path(path).resolve()),
            "sheet": worksheet.title,
            "sheets": names,
            "rows": rows,
            "truncated": truncated,
            "values_are": "формулы" if formulas else "значения, сохранённые "
                          "в файле последней открывавшей его программой",
        }
    finally:
        # read_only workbooks hold the file open until closed, and a lock
        # left on a user's spreadsheet is a support ticket.
        book.close()


def write_xlsx(path: str, rows: Sequence[Sequence[Any]], sheet: str = "",
               overwrite: bool = False) -> Dict[str, Any]:
    """Write rows into a new spreadsheet. Same overwrite rule as write_docx."""
    require("xlsx")
    import openpyxl

    target = Path(path)
    if target.exists() and not overwrite:
        raise FileExistsError(
            f"{target} уже существует; передайте overwrite=True, если заменить "
            f"его — это и есть намерение")

    book = openpyxl.Workbook()
    worksheet = book.active
    if sheet:
        worksheet.title = sheet
    for row in rows:
        worksheet.append(list(row))

    target.parent.mkdir(parents=True, exist_ok=True)
    book.save(str(target))
    return {"path": str(target.resolve()), "rows": len(rows),
            "sheet": worksheet.title, "bytes": target.stat().st_size}
